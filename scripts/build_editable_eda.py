"""Create an editable expanded EDA from verified subset findings and figures."""
from pathlib import Path
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
ASSETS = ROOT / 'output/coco_eda'
summary = json.loads((ASSETS / 'category_summary.json').read_text())
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin = section.bottom_margin = Inches(.6)
section.left_margin = section.right_margin = Inches(.65)
normal = doc.styles['Normal']
normal.font.name = 'Arial'
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.05
for name, size in [('Title',19),('Heading 1',14),('Heading 2',12)]:
    style = doc.styles[name]
    style.font.name = 'Arial'
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0,0,0)
    style.paragraph_format.space_before = Pt(8)
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.keep_with_next = True

def p(text):
    return doc.add_paragraph(text)

def heading(text):
    doc.add_heading(text, level=1)

def image(name, width, caption):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.add_run().add_picture(str(ASSETS/name),width=Inches(width))
    paragraph.paragraph_format.keep_with_next = True
    caption_p = doc.add_paragraph(caption, style='Caption')
    caption_p.paragraph_format.space_after = Pt(6)
    for run in caption_p.runs:
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0,0,0)

def page():
    doc.add_page_break()

doc.add_paragraph('EfficientSAM COCO Exploratory Data Analysis',style='Title')
p('This analysis examines whether a manageable COCO 2017 subset supports a course-scale study of pretrained EfficientSAM. It describes the selected data before preprocessing or training. The findings show substantial category imbalance, many small annotated regions, and an area anomaly that requires measuring combined binary masks rather than relying on stored polygon areas.')
heading('Dataset selection and scope')
p('We selected 800 images from COCO train2017 and 200 from val2017. Image metadata were sorted by image ID, then sampled without replacement using Python random seeds 36219 and 36220, respectively. The selected IDs and filtered annotations were saved for reproducibility. All annotations belonging to each selected image were retained, including crowd entries. The 200 evaluation images remain held out; a development split must be defined within the training data before model selection.')
p('The training subset contains 5,916 annotation entries and the evaluation subset contains 1,642. These are annotation counts, not necessarily counts of individual physical objects: a crowd entry can represent a group. The two subsets share zero COCO image IDs. This checks identity overlap, but does not detect visually similar images with different IDs.')
heading('Training category balance')
image('category_counts.png',5.7,'Figure 1  Ten most frequent categories among the training annotations')
p('Person accounts for 1,869 entries (31.59%), followed by car (313) and chair (252). Annotation frequency therefore varies greatly by category. A single aggregate result could be dominated by common categories and should later be accompanied by category-level results and sample counts.')
p('The subset includes 79 of COCO\'s 80 category definitions. Toothbrush has no training annotations; toaster has one, while microwave and hair drier have three each. These are annotation-coverage gaps, not proof that the corresponding objects are visually absent from every photo.')

page()
heading('Comparison with the source dataset')
image('category_comparison.png',5.6,'Figure 2  Annotation shares in full COCO training and the selected subset')
p('The full training annotation file contains 860,001 entries. Comparing percentages rather than raw counts shows that the subset broadly preserves the shares of the ten most frequent source categories. Person is 30.52% in full COCO and 31.59% in the subset, a difference of 1.07 percentage points; car is 5.10% versus 5.29%, and chair is 4.48% versus 4.26%.')
p('This comparison supports a limited conclusion about common-category balance. It does not establish that rare categories, visual conditions, object sizes, or other datasets are represented adequately. A similar distribution cannot guarantee generalization or rule out later overfitting or underfitting.')
heading('Evaluation coverage and image geometry')
p('The 200 evaluation images contain annotations for 76 categories. Fire hydrant, bear, snowboard, and toaster have no evaluation annotations, so per-category mask-overlap performance cannot be estimated for them in this subset. Results must identify the categories actually evaluated and cannot be presented as full 80-category COCO benchmark results.')
image('aspect_ratios.png',5.0,'Figure 3  Width divided by height for the 800 training images')
p('Training image widths range from 300 to 640 pixels and heights from 200 to 640 pixels. These separate ranges need not describe the same photo. Aspect ratios range from 0.52 to 2.83: 582 images are landscape, 183 portrait, and 35 square. Direct square resizing may distort shapes; the later pipeline must document its resizing policy and transform masks and prompt coordinates consistently.')

page()
heading('Raw images and segmentation challenges')
image('raw_examples.png',6.1,'Figure 4  Original photos and decoded masks for a dog, chair crowd, and bicycle')
p('These examples were deliberately selected to illustrate different challenges. They are not a random sample and should not be used to estimate how frequently those challenges occur. Original photos are shown separately from masks, without overlays or cropping.')
p('The dog in image 458424 occupies a large part of a 640 by 480 photo. Its stored polygon area implies 46.39% coverage; decoded pixel coverage is approximately 46.41%. This small difference is consistent with continuous polygon coordinates and rasterization. Its irregular outline illustrates why segmentation provides more precise localization than a rectangular bounding box.')
p('Image 454495 illustrates crowded seating and occlusion, where people hide parts of chairs. The scattered white regions belong to the chair crowd annotation, not the spectators. In image 252625, the selected bicycle covers only 0.61% of a busy street photo. Thin structures and small regions may lose detail during downsampling.')
for item in summary['visual_sources']:
    cp=p(f"Image {item['id']} source: {item['source']}  License: {item['license']}")
    for run in cp.runs: run.font.size=Pt(8)

page()
heading('Annotation anomaly and data integrity')
image('truck_reference.png',5.8,'Figure 5  Food-truck service window and its combined annotation mask')
p('Truck annotation 397929 in image 475754 has stored area 336,338.72 in a 500 by 500 image, implying 134.54% coverage. Its two polygons remain within image bounds but overlap. Their separate rasterized areas are 241,673 and 94,666 pixels. The union contains 241,673 foreground pixels, corresponding to 96.67% coverage. The stored area closely matches their sum, which counts overlap twice. Source annotations were preserved; subsequent coverage analysis uses combined masks decoded with the COCO toolkit.')
image('mask_coverage.png',5.1,'Figure 6  Coverage measured from decoded training masks including crowd entries')
p('All 5,916 training masks decoded successfully, with zero empty masks. Coverage ranges from 0.0004% to 98.24%, with a median of 0.64%. The histogram is right-skewed: most annotated regions occupy a small fraction of the image, while a smaller number occupy much larger fractions. The first bin represents a range of small values, not empty masks.')
p('All 1,642 evaluation masks also decoded successfully and none were empty. Earlier user-run image checks reported that all 1,000 selected JPEGs opened and matched metadata dimensions, with zero missing or damaged files. Successful decoding establishes technical usability, not semantic correctness of every label. Comprehensive annotation-geometry and visual near-duplicate audits remain outside these completed checks.')
cp=p('Truck image source: '+summary['truck_anomaly']['source']+'  License: https://creativecommons.org/licenses/by-nd/2.0/')
for run in cp.runs: run.font.size=Pt(8)

page()
heading('Evaluation metrics and experimental limits')
p('The primary proposed quality metric is mean foreground Intersection over Union, or IoU, across eligible non-crowd object instances. IoU divides the number of foreground pixels shared by the predicted and reference masks by the number covered by either mask. A perfect match scores one; no overlap scores zero. Averaging object-level scores gives each evaluated instance one contribution rather than letting large masks dominate a global pixel count.')
p('Dice is a companion overlap score. It is twice the shared foreground area divided by the sum of predicted and reference foreground areas. It also ranges from zero to one. Dice and IoU are closely related, so reporting both supports interpretation but does not provide independent evidence of model quality. Both should use the same masks and eligible examples.')
p('Pixel accuracy should not be the primary metric. An object covering 1% of an image could be completely missed while an all-background prediction achieves 99% pixel accuracy. The observed median region coverage of 0.64% makes this concern directly relevant to our subset.')
heading('Planned reporting protocol')
p('Point-prompt and box-prompt results should be reported separately under fixed prompt-generation rules. Results should include category and object-size breakdowns with their sample counts, alongside a macro average over categories with eligible examples. Size definitions must be stated before evaluation: relative image coverage used in this EDA is not interchangeable with COCO\'s absolute small, medium, and large area thresholds.')
p('Crowd entries are useful for understanding data complexity, but should be excluded from the proposed primary single-instance evaluation and counted separately. Categories without eligible evaluation examples must be marked as unavailable rather than assigned an arbitrary score. These prompt-conditioned overlap scores do not reproduce standard COCO mask average precision, which requires a different prediction and matching protocol.')
heading('Implications for the next stage')
p('The pipeline should load images and annotations on demand to suit the available 8 GB RAM, decode instance masks consistently, and verify that geometric operations preserve the alignment of images, masks, and prompts. Image normalization must follow the official EfficientSAM input path to avoid applying it twice. These are requirements for the next implementation stage; preprocessing and model evaluation have not been completed by this EDA.')
p('The study remains a course-scale pilot. Common-category similarity to full COCO does not establish cross-dataset performance. Later experiments should compare the official pretrained baseline and any fine-tuned model on the same held-out examples and hardware, while keeping model-selection data separate. No training curves, accuracy improvements, or runtime claims are inferred from these dataset findings.')
heading('Data and method references')
for text in [
    '[1] COCO Consortium, COCO dataset downloads and annotation documentation. https://cocodataset.org/#download',
    '[2] COCO Consortium, Official Python API and mask toolkit. https://github.com/cocodataset/cocoapi',
    '[3] Y. Xiong et al., EfficientSAM: Leveraged Masked Image Pretraining for Efficient Segment Anything, CVPR 2024. https://arxiv.org/abs/2312.00863',
    'Numerical results were calculated from the saved COCO subset annotations. Image sources and the licenses recorded by COCO appear with the visual examples.'
]:
    cp=p(text)
    for run in cp.runs: run.font.size=Pt(9)

footer=section.footer.paragraphs[0]
footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); footer._p.append(field)
destination=ROOT/'output/docx/EfficientSAM_COCO_EDA_Expanded.docx'
destination.parent.mkdir(parents=True,exist_ok=True)
# Remove inherited decorative paragraph borders from the default template.
for element in [style.element for style in doc.styles] + [paragraph._p for paragraph in doc.paragraphs]:
    for border in list(element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
doc.save(destination)
print(destination)
